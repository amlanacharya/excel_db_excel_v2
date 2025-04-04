"""
Phase 3: Data Recreation
Recreates workbooks from stored data.
"""

import os
import json
from typing import Dict, Any, Optional, List, Set, Tuple

import openpyxl
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from excel_db_excel_v2.phases import Phase
from excel_db_excel_v2.models.workbook import Workbook
from excel_db_excel_v2.models.sheet import Sheet
from excel_db_excel_v2.models.cell import Cell
from excel_db_excel_v2.core.context import ProcessingContext
from excel_db_excel_v2.core.exceptions import PhaseExecutionError
from excel_db_excel_v2.services.database_service import DatabaseService
from excel_db_excel_v2.services.formula_service import FormulaService
from excel_db_excel_v2.services.excel_service import ExcelService
from excel_db_excel_v2.core.utils import ensure_directory_exists


class DataRecreationPhase(Phase):
    """
    Phase 3: Recreate workbooks from stored data.
    Follows Single Responsibility Principle by focusing on just recreation.
    """
    def __init__(self, database_service: DatabaseService, excel_service: ExcelService, 
                 formula_service: FormulaService, config: Any):
        self.database_service = database_service
        self.excel_service = excel_service
        self.formula_service = formula_service
        self.config = config
        self.logger = None
        
    def validate_prerequisites(self, context: ProcessingContext) -> bool:
        """Validate that prerequisites for the phase are met."""
        self.logger = context.logger
        
        if not context.phases_completed.get("data_storage", False):
            self.logger.error("Data storage phase has not been completed")
            context.add_error("data_recreation", "Data storage phase has not been completed")
            return False
            
        if not os.path.exists(self.config.db_filename):
            self.logger.error(f"Database file does not exist: {self.config.db_filename}")
            context.add_error("data_recreation", f"Database file does not exist: {self.config.db_filename}")
            return False
        
        # Ensure output directory exists
        if not ensure_directory_exists(self.config.output_dir):
            self.logger.error(f"Failed to create output directory: {self.config.output_dir}")
            context.add_error("data_recreation", f"Failed to create output directory: {self.config.output_dir}")
            return False
        
        return True
        
    def execute(self, context: ProcessingContext) -> bool:
        """Execute the data recreation phase."""
        try:
            self.logger = context.logger
            self.logger.info("Starting Data Recreation Phase")
            
            # Connect to database
            self.logger.info(f"Connecting to database: {self.config.db_filename}")
            try:
                self.database_service.connect(self.config.db_filename)
            except Exception as e:
                error_msg = f"Failed to connect to database: {str(e)}"
                self.logger.error(error_msg, exc_info=True)
                context.add_error("data_recreation", error_msg, e)
                return False
            
            # Get workbooks from database
            try:
                workbooks = self.database_service.get_workbooks()
                self.logger.info(f"Retrieved {len(workbooks)} workbooks from database")
            except Exception as e:
                error_msg = f"Failed to retrieve workbooks from database: {str(e)}"
                self.logger.error(error_msg, exc_info=True)
                context.add_error("data_recreation", error_msg, e)
                self.database_service.close()
                return False
            
            if not workbooks:
                self.logger.warning("No workbooks found in database")
                context.add_warning("data_recreation", "No workbooks found in database")
                self.database_service.close()
                return False
            
            # Recreate each workbook
            context.recreated_files = []
            success_count = 0
            
            for workbook in workbooks:
                try:
                    self.logger.info(f"Recreating workbook: {workbook.filename}")
                    output_file = self._recreate_workbook(workbook, context)
                    context.recreated_files.append(output_file)
                    success_count += 1
                    self.logger.info(f"Successfully recreated workbook to: {output_file}")
                except Exception as e:
                    error_msg = f"Failed to recreate workbook {workbook.filename}: {str(e)}"
                    self.logger.error(error_msg, exc_info=True)
                    context.add_error("data_recreation", error_msg, e)
                    # Continue with other workbooks
            
            # Close database connection
            self.database_service.close()
            
            if success_count == 0:
                self.logger.warning("No workbooks were successfully recreated")
                context.add_warning("data_recreation", "No workbooks were successfully recreated")
                return False
                
            context.mark_phase_complete("data_recreation")
            self.logger.info(f"Data Recreation Phase completed successfully - {success_count} workbooks recreated")
            return True
            
        except Exception as e:
            self.logger.error(f"Data Recreation Phase failed: {str(e)}", exc_info=True)
            context.add_error("data_recreation", "Failed to recreate workbooks", e)
            if self.database_service.connection:
                self.database_service.close()
            raise PhaseExecutionError("data_recreation", str(e), e)
    
    def _recreate_workbook(self, workbook: Workbook, context: ProcessingContext) -> str:
        """
        Recreate a workbook from the database and return the output file path.
        
        Args:
            workbook: The workbook model to recreate
            context: The processing context
            
        Returns:
            Path to the recreated workbook file
        """
        # Create new openpyxl workbook
        openpyxl_wb = OpenpyxlWorkbook()
        
        # Remove default sheet if necessary
        if len(workbook.sheets) > 0:
            default_sheet = openpyxl_wb.active
            openpyxl_wb.remove(default_sheet)
        
        # Process each sheet
        for sheet_name, sheet in workbook.sheets.items():
            self.logger.info(f"Recreating sheet: {sheet_name}")
            
            # Create the sheet
            ws = openpyxl_wb.create_sheet(title=sheet_name)
            
            # Apply merged cells
            if sheet.merged_cells:
                self.logger.debug(f"Applying {len(sheet.merged_cells)} merged cell ranges")
                for merged_range in sheet.merged_cells:
                    ws.merge_cells(merged_range)
            
            # Apply dimensions
            if sheet.column_dimensions:
                self.logger.debug(f"Applying {len(sheet.column_dimensions)} column dimensions")
                for col_key, properties in sheet.column_dimensions.items():
                    if col_key in ws.column_dimensions and properties.get("width"):
                        ws.column_dimensions[col_key].width = properties["width"]
            
            if sheet.row_dimensions:
                self.logger.debug(f"Applying {len(sheet.row_dimensions)} row dimensions")
                for row_key, properties in sheet.row_dimensions.items():
                    try:
                        row = int(row_key)
                        if row in ws.row_dimensions and properties.get("height"):
                            ws.row_dimensions[row].height = properties["height"]
                    except (ValueError, TypeError):
                        self.logger.warning(f"Invalid row key: {row_key}")
            
            # Populate cells
            self.logger.debug(f"Populating {len(sheet.cells)} cells")
            for coordinate, cell in sheet.cells.items():
                self._set_cell_value(ws, coordinate, cell, sheet_name)
                
            self.logger.info(f"Sheet {sheet_name} recreated with {len(sheet.cells)} cells")
        
        # Add workbook links for formulas if needed
        if self.config.include_links_sheet and 'Form X Report' in workbook.filename:
            self.logger.info("Adding _Links sheet with workbook references")
            links_sheet = openpyxl_wb.create_sheet(title="_Links", index=0)
            links_sheet["A1"] = "Workbook Index References"
            links_sheet["A2"] = "[1] = Deposits Data Lite.xlsx"
            links_sheet["A3"] = "[2] = Loans Data Lite.xlsx"
            links_sheet["A4"] = "[3] = Form X Report  Main Lite.xlsx"
            links_sheet["A6"] = "Note: These links help resolve formulas with [1], [2] references."
            links_sheet["A7"] = "You may need to update links manually in Excel: Data > Edit Links"
        
        # Handle external links
        if self.config.fix_external_references and self.config.new_base_path:
            self.logger.info(f"Updating external links with new base path: {self.config.new_base_path}")
            links_result = self.excel_service.update_external_links(openpyxl_wb, self.config.new_base_path)
            self.logger.debug(f"External links update result: {links_result}")
        
        # Save the workbook
        base_name = os.path.splitext(workbook.filename)[0]
        output_file = os.path.join(self.config.output_dir, f"{base_name}{self.config.output_suffix}.xlsx")
        
        try:
            openpyxl_wb.save(output_file)
            self.logger.info(f"Saved workbook to {output_file}")
            return output_file
        except Exception as e:
            error_msg = f"Failed to save workbook to {output_file}: {str(e)}"
            self.logger.error(error_msg, exc_info=True)
            raise Exception(error_msg) from e
    
    def _set_cell_value(self, worksheet, coordinate, cell, sheet_name):
        """
        Set the value of a cell in the worksheet.
        
        Args:
            worksheet: The openpyxl worksheet
            coordinate: The cell coordinate
            cell: The cell model
            sheet_name: The name of the sheet (for logging)
        """
        try:
            if cell.is_formula:
                # Handle formula cells
                formula_value = cell.value
                
                # Special handling for indexed references in specific sheets
                is_special_sheet = sheet_name in ["MIS-Report", "Part I", "Part II", "Part III"]
                has_indexed_ref = ('[1]' in formula_value) or ('[2]' in formula_value) or ('[3]' in formula_value)
                
                try:
                    if is_special_sheet and has_indexed_ref:
                        # For indexed references, set value directly
                        worksheet[coordinate].value = formula_value
                    elif formula_value.startswith('='):
                        # Standard formula
                        worksheet[coordinate].value = None
                        worksheet[coordinate].formula = formula_value[1:]
                    else:
                        # Non-standard formula
                        worksheet[coordinate].value = None
                        worksheet[coordinate].formula = formula_value
                except Exception as e:
                    self.logger.warning(f"Error setting formula in {sheet_name}!{coordinate}: {e}")
                    worksheet[coordinate].value = formula_value
            else:
                # Handle non-formula cells
                value = cell.value
                
                # Convert value to appropriate type
                if isinstance(value, str):
                    if value.lower() == 'true':
                        worksheet[coordinate] = True
                    elif value.lower() == 'false':
                        worksheet[coordinate] = False
                    else:
                        try:
                            if value.isdigit():
                                worksheet[coordinate] = int(value)
                            else:
                                try:
                                    worksheet[coordinate] = float(value)
                                except ValueError:
                                    worksheet[coordinate] = value
                        except (ValueError, TypeError, AttributeError):
                            worksheet[coordinate] = value
                else:
                    worksheet[coordinate] = value
                
        except Exception as e:
            self.logger.warning(f"Error setting value in {sheet_name}!{coordinate}: {e}")
            try:
                # Fallback to setting as string
                worksheet[coordinate] = str(cell.value)
            except Exception:
                self.logger.error(f"Failed to set value in {sheet_name}!{coordinate} even as string")