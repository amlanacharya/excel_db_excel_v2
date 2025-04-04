"""
Phase 4: Font Color Correction
Fixes font colors in recreated workbooks.
"""

import os
import json
from typing import Dict, Any, Optional, List, Set, Tuple
from copy import copy

import openpyxl
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from excel_db_excel_v2.phases import Phase
from excel_db_excel_v2.core.context import ProcessingContext
from excel_db_excel_v2.core.exceptions import PhaseExecutionError
from excel_db_excel_v2.core.utils import ensure_directory_exists
from excel_db_excel_v2.services.excel_service import ExcelService


class FontColorCorrectionPhase(Phase):
    """
    Phase 4: Fix font colors in recreated workbooks.
    Follows Single Responsibility Principle by focusing on just color correction.
    """
    def __init__(self, excel_service: ExcelService, config: Any):
        self.excel_service = excel_service
        self.config = config
        self.logger = None
        
    def validate_prerequisites(self, context: ProcessingContext) -> bool:
        """Validate that prerequisites for the phase are met."""
        self.logger = context.logger
        
        if not context.phases_completed.get("data_recreation", False):
            self.logger.error("Data recreation phase has not been completed")
            context.add_error("font_color_correction", "Data recreation phase has not been completed")
            return False
            
        if not context.recreated_files:
            self.logger.error("No recreated files to process")
            context.add_error("font_color_correction", "No recreated files to process")
            return False
        
        missing_files = []
        for file in context.recreated_files:
            if not os.path.exists(file):
                missing_files.append(file)
                
        if missing_files:
            self.logger.warning(f"Some recreated files are missing: {missing_files}")
            context.add_warning("font_color_correction", 
                               f"Some recreated files are missing: {missing_files}")
            
            if len(missing_files) == len(context.recreated_files):
                self.logger.error("All recreated files are missing")
                context.add_error("font_color_correction", "All recreated files are missing")
                return False
        
        return True
        
    def execute(self, context: ProcessingContext) -> bool:
        """Execute the font color correction phase."""
        try:
            self.logger = context.logger
            self.logger.info("Starting Font Color Correction Phase")
            
            # Ensure output directory exists
            if not ensure_directory_exists(self.config.output_dir):
                self.logger.error(f"Failed to create output directory: {self.config.output_dir}")
                context.add_error("font_color_correction", 
                                 f"Failed to create output directory: {self.config.output_dir}")
                return False
            
            # Process each file
            context.fixed_files = []
            success_count = 0
            
            for file in context.recreated_files:
                if not os.path.exists(file):
                    self.logger.warning(f"Skipping missing file: {file}")
                    context.add_warning("font_color_correction", f"Skipping missing file: {file}")
                    continue
                    
                try:
                    self.logger.info(f"Processing file: {file}")
                    output_file = self._fix_font_colors(file, context)
                    context.fixed_files.append(output_file)
                    success_count += 1
                    self.logger.info(f"Successfully processed file: {output_file}")
                except Exception as e:
                    error_msg = f"Failed to process file {file}: {str(e)}"
                    self.logger.error(error_msg, exc_info=True)
                    context.add_error("font_color_correction", error_msg, e)
                    # Continue with other files
            
            if success_count == 0:
                self.logger.warning("No files were successfully processed")
                context.add_warning("font_color_correction", "No files were successfully processed")
                return False
                
            context.mark_phase_complete("font_color_correction")
            self.logger.info(f"Font Color Correction Phase completed successfully - {success_count} files processed")
            return True
            
        except Exception as e:
            self.logger.error(f"Font Color Correction Phase failed: {str(e)}", exc_info=True)
            context.add_error("font_color_correction", "Failed to fix font colors", e)
            raise PhaseExecutionError("font_color_correction", str(e), e)
    
    def _fix_font_colors(self, file: str, context: ProcessingContext) -> str:
        """
        Fix font colors in a workbook and return the output file path.
        
        Args:
            file: Path to the workbook file
            context: The processing context
            
        Returns:
            Path to the fixed workbook file
        """
        # Load the workbook
        try:
            workbook = openpyxl.load_workbook(file)
        except Exception as e:
            error_msg = f"Failed to load workbook {file}: {str(e)}"
            self.logger.error(error_msg, exc_info=True)
            raise Exception(error_msg) from e
        
        total_cells_modified = 0
        
        # Process each sheet
        for sheet_name in workbook.sheetnames:
            self.logger.info(f"Processing sheet: {sheet_name}")
            ws = workbook[sheet_name]
            
            # Process all cells
            cells_modified = 0
            cells_processed = 0
            
            for row in range(1, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row, column=col)
                    cells_processed += 1
                    
                    # Skip empty cells
                    if cell.value is None:
                        continue
                    
                    # Fix font color
                    if self._fix_cell_font_color(cell):
                        cells_modified += 1
            
            self.logger.info(f"Modified {cells_modified}/{cells_processed} cells in sheet {sheet_name}")
            total_cells_modified += cells_modified
        
        # Create output filename with the appropriate suffix
        base_name, ext = os.path.splitext(file)
        if self.config.output_suffix:
            # If the input file already has a suffix (like _recreated), replace it
            if "_" in os.path.basename(base_name):
                parts = base_name.split("_")
                base_name = "_".join(parts[:-1])
            
            output_file = f"{base_name}{self.config.output_suffix}{ext}"
        else:
            # If no suffix specified, append _fixed
            output_file = f"{base_name}_fixed{ext}"
        
        # Save the modified workbook
        try:
            workbook.save(output_file)
            self.logger.info(f"Saved fixed workbook to {output_file} (modified {total_cells_modified} cells)")
            return output_file
        except Exception as e:
            error_msg = f"Failed to save workbook to {output_file}: {str(e)}"
            self.logger.error(error_msg, exc_info=True)
            raise Exception(error_msg) from e
    
    def _fix_cell_font_color(self, cell) -> bool:
        """
        Fix the font color of a cell.
        
        Args:
            cell: The openpyxl cell
            
        Returns:
            True if the cell was modified, False otherwise
        """
        try:
            # Get current font
            current_font = cell.font
            
            # Check if we need to change the color
            if not hasattr(current_font, 'color') or current_font.color is None or current_font.color.rgb == self.config.target_font_color:
                return False
            
            # Make a copy of the current font with the target color
            new_font = copy(current_font)
            new_font.color = self.config.target_font_color
            
            # Set the new font
            cell.font = new_font
            
            return True
        except Exception as e:
            self.logger.warning(f"Error fixing font in cell {cell.coordinate}: {e}")
            return False