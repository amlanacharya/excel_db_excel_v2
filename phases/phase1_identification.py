"""
Phase 1: Data Identification
Identifies and extracts data from Excel files.
"""

import os
import json
from datetime import datetime,date
from typing import Dict, Any, Optional, List, Set, Tuple

import openpyxl
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from excel_db_excel_v2.phases import Phase
from excel_db_excel_v2.models.workbook import Workbook
from excel_db_excel_v2.models.sheet import Sheet
from excel_db_excel_v2.models.cell import Cell
from excel_db_excel_v2.core.context import ProcessingContext
from excel_db_excel_v2.core.exceptions import PhaseExecutionError
from excel_db_excel_v2.core.utils import serialize_to_json, DateTimeEncoder
from excel_db_excel_v2.services.excel_service import ExcelService


class DataIdentificationPhase(Phase):
    """
    Phase 1: Identify and extract data from Excel files.
    Follows Single Responsibility Principle by focusing on just identification.
    """
    def __init__(self, excel_service: ExcelService, config: Any):
        self.excel_service = excel_service
        self.config = config
        self.logger = None
        
    def validate_prerequisites(self, context: ProcessingContext) -> bool:
        """Validate that prerequisites for the phase are met."""
        self.logger = context.logger
        
        if not self.config.excel_files:
            self.logger.error("No Excel files specified for processing")
            context.add_error("data_identification", "No Excel files specified for processing")
            return False
        
        missing_files = []
        for file in self.config.excel_files:
            if not os.path.exists(file):
                missing_files.append(file)
                context.add_warning("data_identification", f"Input file '{file}' does not exist")
        
        if missing_files and len(missing_files) == len(self.config.excel_files):
            self.logger.error("All specified Excel files are missing")
            context.add_error("data_identification", "All specified Excel files are missing")
            return False
        
        return True
        
    def execute(self, context: ProcessingContext) -> bool:
        """Execute the data identification phase."""
        try:
            self.logger = context.logger
            self.logger.info("Starting Data Identification Phase")
            
            # Reset workbook data
            context.workbook_data = {}
            
            # Process each Excel file
            for file in self.config.excel_files:
                if not os.path.exists(file):
                    self.logger.warning(f"Skipping non-existent file: {file}")
                    context.add_warning("data_identification", f"Skipping non-existent file: {file}")
                    continue
                    
                self.logger.info(f"Processing file: {file}")
                workbook = self._process_file(file, context)
                context.workbook_data[file] = workbook.to_dict()
            
            # Save identification results to JSON if configured
            if self.config.save_identification_json:
                self.logger.info(f"Saving identification results to {self.config.identification_file}")
                if serialize_to_json(context.workbook_data, self.config.identification_file):
                    self.logger.info("Identification results saved successfully")
                else:
                    context.add_warning("data_identification", f"Failed to save identification results to {self.config.identification_file}")
            
            context.mark_phase_complete("data_identification")
            self.logger.info("Data Identification Phase completed successfully")
            return True
            
        except Exception as e:
            self.logger.error(f"Data Identification Phase failed: {str(e)}", exc_info=True)
            context.add_error("data_identification", "Failed to identify data", e)
            raise PhaseExecutionError("data_identification", str(e), e)
    
    def _process_file(self, file: str, context: ProcessingContext) -> Workbook:
        """
        Process a single Excel file and return a Workbook model.
        
        Args:
            file: Path to the Excel file
            context: Processing context
            
        Returns:
            Workbook model with data extracted from the file
        """
        self.logger.info(f"Loading workbook: {file}")
        
        # Load the workbook with openpyxl
        try:
            openpyxl_wb = openpyxl.load_workbook(file, data_only=False)
        except Exception as e:
            error_msg = f"Failed to load workbook {file}: {str(e)}"
            self.logger.error(error_msg)
            context.add_error("data_identification", error_msg, e)
            raise PhaseExecutionError("data_identification", error_msg, e)
        
        # Extract workbook metadata
        if self.config.extract_metadata:
            properties = self._extract_workbook_metadata(openpyxl_wb)
        else:
            properties = {'sheet_names': openpyxl_wb.sheetnames}
        
        # Create the workbook model
        workbook = Workbook(
            filename=os.path.basename(file),
            properties=properties
        )
        
        # Process each sheet in the workbook
        for sheet_name in openpyxl_wb.sheetnames:
            # Skip excluded sheets
            if sheet_name in self.config.exclude_sheets:
                self.logger.info(f"Skipping excluded sheet: {sheet_name}")
                continue
            
            # Get the worksheet
            ws = openpyxl_wb[sheet_name]
            
            # Create the sheet model
            sheet_type = "report" if sheet_name in self.config.report_sheets else "non_report"
            sheet = Sheet(
                name=sheet_name,
                sheet_type=sheet_type,
                max_row=ws.max_row,
                max_column=ws.max_column
            )
            
            # Add merged cells
            sheet.merged_cells = [str(merged_range) for merged_range in ws.merged_cells.ranges]
            
            # Add column dimensions
            sheet.column_dimensions = {
                col: {"width": ws.column_dimensions[col].width} 
                for col in ws.column_dimensions
            }
            
            # Add row dimensions
            sheet.row_dimensions = {
                row: {"height": ws.row_dimensions[row].height} 
                for row in ws.row_dimensions
            }
            
            # Process all cells in the sheet
            self.logger.debug(f"Processing cells in sheet: {sheet_name}")
            cells_processed = 0
            formulas_found = 0
            references_found = 0
            
            for row in range(1, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    openpyxl_cell = ws.cell(row=row, column=col)
                    
                    # Skip empty cells
                    if openpyxl_cell.value is None:
                        continue
                    
                    # Create the cell model
                    cell = self._process_cell(openpyxl_cell)
                    sheet.add_cell(cell)
                    
                    cells_processed += 1
                    if cell.is_formula:
                        formulas_found += 1
                    if self._has_external_reference(cell.value):
                        references_found += 1
            
            # Add the sheet to the workbook
            workbook.add_sheet(sheet)
            
            self.logger.info(f"Processed sheet '{sheet_name}': {cells_processed} cells, {formulas_found} formulas, {references_found} external references")
        
        self.logger.info(f"Completed processing file: {file}")
        return workbook
    
    def _extract_workbook_metadata(self, openpyxl_wb: OpenpyxlWorkbook) -> Dict[str, Any]:
        """
        Extract metadata from an openpyxl workbook.
        
        Args:
            openpyxl_wb: The openpyxl workbook
            
        Returns:
            Dictionary with workbook metadata
        """
        properties = {
            'title': openpyxl_wb.properties.title,
            'creator': openpyxl_wb.properties.creator,
            'created': str(openpyxl_wb.properties.created) if openpyxl_wb.properties.created else None,
            'sheet_names': openpyxl_wb.sheetnames
        }
        
        # Add more metadata if available
        if hasattr(openpyxl_wb.properties, 'modified'):
            properties['modified'] = str(openpyxl_wb.properties.modified) if openpyxl_wb.properties.modified else None
            
        if hasattr(openpyxl_wb.properties, 'lastModifiedBy'):
            properties['last_modified_by'] = openpyxl_wb.properties.lastModifiedBy
            
        return properties
    
    def _process_cell(self, openpyxl_cell) -> Cell:
        """
        Process an openpyxl cell and create a Cell model.
        
        Args:
            openpyxl_cell: The openpyxl cell
            
        Returns:
            Cell model
        """
        coordinate = openpyxl_cell.coordinate
        value = openpyxl_cell.value
        
        # Determine if the cell contains a formula
        is_formula = False
        if isinstance(value, str) and value.startswith('=') and self.config.parse_formulas:
            is_formula = True
        
        # Handle datetime objects
        if isinstance(value, (datetime, date)):
            value = value.isoformat()
        
        return Cell(
            coordinate=coordinate,
            value=value,
            is_formula=is_formula
        )
    
    def _has_external_reference(self, cell_value) -> bool:
        """
        Check if a cell value contains external references.
        
        Args:
            cell_value: The cell value to check
            
        Returns:
            True if the value contains external references, False otherwise
        """
        if not isinstance(cell_value, str) or not self.config.detect_references:
            return False
        
        # Check for external references
        return (
            '.xlsx' in cell_value or 
            '.xls' in cell_value or 
            ('[' in cell_value and ']' in cell_value)
        )